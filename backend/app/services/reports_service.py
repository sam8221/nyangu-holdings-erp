"""Business reports. Each report returns columns + rows + totals, exportable to CSV or Excel."""

from __future__ import annotations

import csv
import io
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.assets import Asset
from app.models.finance import BillStatus, Expense, ExpenseCategory, ExpenseStatus, SupplierBill
from app.models.hr import Employee, EmployeeStatus
from app.models.inventory import StockLevel, StockMovement
from app.models.organization import Branch, Department
from app.models.partners import Customer, Product, ProductCategory
from app.models.procurement import POStatus, PurchaseOrder
from app.models.sales import InvoiceStatus, SalesInvoice, SalesInvoiceLine
from app.services.finance_service import FinanceService
from app.utils.exceptions import BadRequestError, NotFoundError
from app.utils.money import ZERO, money, qty
from app.utils.time import local_today, start_of_local_day


@dataclass
class Column:
    key: str
    label: str
    kind: str = "text"  # text | money | quantity | number | percent | date


@dataclass
class ReportResult:
    key: str
    title: str
    columns: list[Column]
    rows: list[dict[str, Any]]
    totals: dict[str, Any] | None = None
    parameters: dict[str, Any] = field(default_factory=dict)

    def as_dict(self) -> dict[str, Any]:
        return {
            "key": self.key,
            "title": self.title,
            "columns": [c.__dict__ for c in self.columns],
            "rows": self.rows,
            "totals": self.totals,
            "parameters": self.parameters,
        }


@dataclass(frozen=True)
class ReportParams:
    date_from: date
    date_to: date
    group_by: str | None = None


@dataclass(frozen=True)
class ReportDefinition:
    key: str
    title: str
    description: str
    uses_dates: bool
    builder: Callable[[Session, ReportParams], ReportResult]
    group_by_options: tuple[str, ...] = ()


def _sum(rows: list[dict[str, Any]], *keys: str) -> dict[str, Any]:
    return {k: sum((r[k] for r in rows), ZERO) for k in keys}


POSTED = InvoiceStatus.POSTED


# ---------------------------------------------------------------------- sales
def sales_summary(db: Session, p: ReportParams) -> ReportResult:
    group = p.group_by or "month"
    if group == "customer":
        key_expr = Customer.name
        label = "Customer"
    elif group == "day":
        key_expr = func.to_char(SalesInvoice.invoice_date, "YYYY-MM-DD")
        label = "Date"
    else:
        key_expr = func.to_char(SalesInvoice.invoice_date, "YYYY-MM")
        label = "Month"
    stmt = (
        select(
            key_expr.label("period"),
            func.count(SalesInvoice.id),
            func.sum(SalesInvoice.subtotal),
            func.sum(SalesInvoice.tax_total),
            func.sum(SalesInvoice.total),
            func.sum(SalesInvoice.amount_paid),
        )
        .join(Customer, Customer.id == SalesInvoice.customer_id)
        .where(
            SalesInvoice.status.in_(POSTED),
            SalesInvoice.invoice_date >= p.date_from,
            SalesInvoice.invoice_date <= p.date_to,
        )
        .group_by(key_expr)
        .order_by(key_expr)
    )
    rows = [
        {
            "group": g,
            "invoices": n,
            "net_sales": money(sub),
            "vat": money(tax),
            "gross_sales": money(total),
            "collected": money(paid),
        }
        for g, n, sub, tax, total, paid in db.execute(stmt).all()
    ]
    totals = {"invoices": sum(r["invoices"] for r in rows)} | _sum(
        rows, "net_sales", "vat", "gross_sales", "collected"
    )
    return ReportResult(
        "sales-summary",
        "Sales summary",
        [
            Column("group", label),
            Column("invoices", "Invoices", "number"),
            Column("net_sales", "Net sales", "money"),
            Column("vat", "VAT", "money"),
            Column("gross_sales", "Gross sales", "money"),
            Column("collected", "Collected to date", "money"),
        ],
        rows,
        totals,
    )


def sales_by_product(db: Session, p: ReportParams) -> ReportResult:
    cost = func.sum(func.round(SalesInvoiceLine.quantity * SalesInvoiceLine.unit_cost, 2))
    stmt = (
        select(
            Product.sku,
            Product.name,
            func.sum(SalesInvoiceLine.quantity),
            func.sum(SalesInvoiceLine.line_subtotal),
            cost,
        )
        .join(SalesInvoice, SalesInvoice.id == SalesInvoiceLine.invoice_id)
        .join(Product, Product.id == SalesInvoiceLine.product_id)
        .where(
            SalesInvoice.status.in_(POSTED),
            SalesInvoice.invoice_date >= p.date_from,
            SalesInvoice.invoice_date <= p.date_to,
        )
        .group_by(Product.sku, Product.name)
        .order_by(func.sum(SalesInvoiceLine.line_subtotal).desc())
    )
    rows = []
    for sku, name, quantity, revenue, cogs in db.execute(stmt).all():
        revenue, cogs = money(revenue), money(cogs or 0)
        profit = revenue - cogs
        rows.append(
            {
                "sku": sku,
                "name": name,
                "quantity": qty(quantity),
                "revenue": revenue,
                "cost_of_sales": cogs,
                "gross_profit": profit,
                "margin_percent": money(profit * 100 / revenue) if revenue else ZERO,
            }
        )
    totals = _sum(rows, "revenue", "cost_of_sales", "gross_profit")
    totals["margin_percent"] = (
        money(totals["gross_profit"] * 100 / totals["revenue"]) if totals["revenue"] else ZERO
    )
    return ReportResult(
        "sales-by-product",
        "Sales and gross profit by product",
        [
            Column("sku", "SKU"),
            Column("name", "Product"),
            Column("quantity", "Quantity", "quantity"),
            Column("revenue", "Net sales", "money"),
            Column("cost_of_sales", "Cost of sales", "money"),
            Column("gross_profit", "Gross profit", "money"),
            Column("margin_percent", "Margin %", "percent"),
        ],
        rows,
        totals,
    )


# ---------------------------------------------------------------------- purchasing & expenses
def purchases_by_supplier(db: Session, p: ReportParams) -> ReportResult:
    orders = db.scalars(
        select(PurchaseOrder).where(
            PurchaseOrder.status.notin_((POStatus.DRAFT, POStatus.CANCELLED)),
            PurchaseOrder.order_date >= p.date_from,
            PurchaseOrder.order_date <= p.date_to,
        )
    ).all()
    per: dict[str, dict[str, Any]] = {}
    for order in orders:
        row = per.setdefault(
            order.supplier.name,
            {"supplier": order.supplier.name, "orders": 0, "ordered": ZERO, "received": ZERO},
        )
        row["orders"] += 1
        row["ordered"] += order.subtotal
        row["received"] += order.received_value
    rows = sorted(per.values(), key=lambda r: r["ordered"], reverse=True)
    totals = {"orders": sum(r["orders"] for r in rows)} | _sum(rows, "ordered", "received")
    return ReportResult(
        "purchases-by-supplier",
        "Purchases by supplier",
        [
            Column("supplier", "Supplier"),
            Column("orders", "Orders", "number"),
            Column("ordered", "Ordered (net)", "money"),
            Column("received", "Received (net)", "money"),
        ],
        rows,
        totals,
    )


def expenses_by_category(db: Session, p: ReportParams) -> ReportResult:
    stmt = (
        select(
            ExpenseCategory.name,
            func.count(Expense.id),
            func.sum(Expense.amount),
            func.sum(Expense.tax_amount),
        )
        .join(ExpenseCategory, ExpenseCategory.id == Expense.category_id)
        .where(
            Expense.status == ExpenseStatus.APPROVED,
            Expense.expense_date >= p.date_from,
            Expense.expense_date <= p.date_to,
        )
        .group_by(ExpenseCategory.name)
        .order_by(func.sum(Expense.amount).desc())
    )
    rows = [
        {"category": name, "count": n, "amount": money(a), "vat": money(t), "total": money(a + t)}
        for name, n, a, t in db.execute(stmt).all()
    ]
    totals = {"count": sum(r["count"] for r in rows)} | _sum(rows, "amount", "vat", "total")
    return ReportResult(
        "expenses-by-category",
        "Approved expenses by category",
        [
            Column("category", "Category"),
            Column("count", "Expenses", "number"),
            Column("amount", "Amount", "money"),
            Column("vat", "VAT", "money"),
            Column("total", "Total", "money"),
        ],
        rows,
        totals,
    )


# ---------------------------------------------------------------------- profit, VAT
def profit_and_loss(db: Session, p: ReportParams) -> ReportResult:
    period = (SalesInvoice.invoice_date >= p.date_from, SalesInvoice.invoice_date <= p.date_to)
    revenue = money(
        db.scalar(
            select(func.coalesce(func.sum(SalesInvoice.subtotal), 0)).where(
                SalesInvoice.status.in_(POSTED), *period
            )
        )
    )
    cogs = money(
        db.scalar(
            select(
                func.coalesce(
                    func.sum(func.round(SalesInvoiceLine.quantity * SalesInvoiceLine.unit_cost, 2)),
                    0,
                )
            )
            .join(SalesInvoice, SalesInvoice.id == SalesInvoiceLine.invoice_id)
            .where(SalesInvoice.status.in_(POSTED), *period)
        )
    )
    expense_rows = db.execute(
        select(ExpenseCategory.name, func.sum(Expense.amount))
        .join(ExpenseCategory, ExpenseCategory.id == Expense.category_id)
        .where(
            Expense.status == ExpenseStatus.APPROVED,
            Expense.expense_date >= p.date_from,
            Expense.expense_date <= p.date_to,
        )
        .group_by(ExpenseCategory.name)
        .order_by(ExpenseCategory.name)
    ).all()
    gross = revenue - cogs
    rows: list[dict[str, Any]] = [
        {"line": "Revenue (net of VAT)", "amount": revenue},
        {"line": "Cost of sales", "amount": -cogs},
        {"line": "Gross profit", "amount": gross},
    ]
    operating = ZERO
    for name, amount in expense_rows:
        operating += money(amount)
        rows.append({"line": f"Expense: {name}", "amount": -money(amount)})
    rows.append({"line": "Total operating expenses", "amount": -operating})
    net = gross - operating
    rows.append({"line": "Net profit", "amount": net})
    return ReportResult(
        "profit-and-loss",
        "Profit and loss",
        [Column("line", "Line"), Column("amount", "Amount", "money")],
        rows,
        {"amount": net},
    )


def vat_summary(db: Session, p: ReportParams) -> ReportResult:
    output_vat = money(
        db.scalar(
            select(func.coalesce(func.sum(SalesInvoice.tax_total), 0)).where(
                SalesInvoice.status.in_(POSTED),
                SalesInvoice.invoice_date >= p.date_from,
                SalesInvoice.invoice_date <= p.date_to,
            )
        )
    )
    bill_vat = money(
        db.scalar(
            select(func.coalesce(func.sum(SupplierBill.tax_total), 0)).where(
                SupplierBill.status.notin_((BillStatus.PENDING, BillStatus.CANCELLED)),
                SupplierBill.bill_date >= p.date_from,
                SupplierBill.bill_date <= p.date_to,
            )
        )
    )
    expense_vat = money(
        db.scalar(
            select(func.coalesce(func.sum(Expense.tax_amount), 0)).where(
                Expense.status == ExpenseStatus.APPROVED,
                Expense.expense_date >= p.date_from,
                Expense.expense_date <= p.date_to,
            )
        )
    )
    net = output_vat - bill_vat - expense_vat
    rows = [
        {"line": "Output VAT on sales", "amount": output_vat},
        {"line": "Input VAT on supplier bills", "amount": -bill_vat},
        {"line": "Input VAT on expenses", "amount": -expense_vat},
        {"line": "Net VAT payable (refundable if negative)", "amount": net},
    ]
    return ReportResult(
        "vat-summary",
        "VAT summary",
        [Column("line", "Line"), Column("amount", "Amount", "money")],
        rows,
        {"amount": net},
    )


# ---------------------------------------------------------------------- stock
def inventory_valuation(db: Session, p: ReportParams) -> ReportResult:
    totals_sq = (
        select(StockLevel.product_id, func.sum(StockLevel.quantity).label("qty"))
        .group_by(StockLevel.product_id)
        .subquery()
    )
    stmt = (
        select(Product, ProductCategory.name, totals_sq.c.qty)
        .join(totals_sq, totals_sq.c.product_id == Product.id)
        .outerjoin(ProductCategory, ProductCategory.id == Product.category_id)
        .where(totals_sq.c.qty > 0)
        .order_by(Product.sku)
    )
    rows = []
    for product, category, quantity in db.execute(stmt).all():
        quantity = qty(quantity)
        rows.append(
            {
                "sku": product.sku,
                "name": product.name,
                "category": category or "",
                "quantity": quantity,
                "unit_cost": product.cost_price,
                "value": money(quantity * product.cost_price),
                "retail_value": money(quantity * product.selling_price),
            }
        )
    return ReportResult(
        "inventory-valuation",
        "Inventory valuation (weighted average cost)",
        [
            Column("sku", "SKU"),
            Column("name", "Product"),
            Column("category", "Category"),
            Column("quantity", "On hand", "quantity"),
            Column("unit_cost", "Average cost", "money"),
            Column("value", "Stock value", "money"),
            Column("retail_value", "Retail value", "money"),
        ],
        rows,
        _sum(rows, "value", "retail_value"),
        {"as_of": local_today()},
    )


def stock_movement_summary(db: Session, p: ReportParams) -> ReportResult:
    start = start_of_local_day(p.date_from)
    end = start_of_local_day(date.fromordinal(p.date_to.toordinal() + 1))
    inbound = func.sum(func.greatest(StockMovement.quantity, 0))
    outbound = func.sum(func.least(StockMovement.quantity, 0))
    stmt = (
        select(Product.sku, Product.name, inbound, outbound, func.sum(StockMovement.quantity))
        .join(Product, Product.id == StockMovement.product_id)
        .where(StockMovement.created_at >= start, StockMovement.created_at < end)
        .group_by(Product.sku, Product.name)
        .order_by(Product.sku)
    )
    rows = [
        {"sku": sku, "name": name, "in": qty(i), "out": qty(-o), "net": qty(n)}
        for sku, name, i, o, n in db.execute(stmt).all()
    ]
    return ReportResult(
        "stock-movements",
        "Stock movement summary",
        [
            Column("sku", "SKU"),
            Column("name", "Product"),
            Column("in", "Quantity in", "quantity"),
            Column("out", "Quantity out", "quantity"),
            Column("net", "Net change", "quantity"),
        ],
        rows,
        None,
    )


# ---------------------------------------------------------------------- aging, people, assets
def _aging_result(key: str, title: str, party_label: str, data: dict[str, Any]) -> ReportResult:
    buckets = ("current", "days_1_30", "days_31_60", "days_61_90", "days_over_90", "total")
    rows = [{"party": r["party"].name, **{b: r[b] for b in buckets}} for r in data["rows"]]
    labels = ("Current", "1-30 days", "31-60 days", "61-90 days", "Over 90 days", "Total")
    return ReportResult(
        key,
        title,
        [Column("party", party_label)]
        + [Column(b, lbl, "money") for b, lbl in zip(buckets, labels, strict=True)],
        rows,
        data["totals"],
        {"as_of": data["as_of"]},
    )


def receivables_aging(db: Session, p: ReportParams) -> ReportResult:
    return _aging_result(
        "receivables-aging", "Receivables aging", "Customer", FinanceService(db).receivables_aging()
    )


def payables_aging(db: Session, p: ReportParams) -> ReportResult:
    return _aging_result(
        "payables-aging", "Payables aging", "Supplier", FinanceService(db).payables_aging()
    )


def headcount(db: Session, p: ReportParams) -> ReportResult:
    group = p.group_by or "department"
    if group == "branch":
        key_expr, label = func.coalesce(Branch.name, "(no branch)"), "Branch"
        stmt = select(key_expr, Employee.status, func.count()).outerjoin(
            Branch, Branch.id == Employee.branch_id
        )
    else:
        key_expr, label = func.coalesce(Department.name, "(no department)"), "Department"
        stmt = select(key_expr, Employee.status, func.count()).outerjoin(
            Department, Department.id == Employee.department_id
        )
    stmt = stmt.group_by(key_expr, Employee.status)
    per: dict[str, dict[str, Any]] = {}
    for name, status, n in db.execute(stmt).all():
        row = per.setdefault(name, {"group": name, **dict.fromkeys(EmployeeStatus.ALL, 0)})
        row[status] = n
    rows = sorted(per.values(), key=lambda r: r["group"])
    for row in rows:
        row["total"] = sum(row[s] for s in EmployeeStatus.ALL)
    totals = {k: sum(r[k] for r in rows) for k in (*EmployeeStatus.ALL, "total")}
    return ReportResult(
        "headcount",
        "Employee headcount",
        [Column("group", label)]
        + [Column(s, s.replace("_", " ").title(), "number") for s in EmployeeStatus.ALL]
        + [Column("total", "Total", "number")],
        rows,
        totals,
    )


def asset_register(db: Session, p: ReportParams) -> ReportResult:
    assets = db.scalars(select(Asset).order_by(Asset.asset_tag)).all()
    rows = [
        {
            "asset_tag": a.asset_tag,
            "name": a.name,
            "category": a.category.name,
            "status": a.status,
            "purchase_date": a.purchase_date,
            "cost": a.purchase_cost,
            "accumulated_depreciation": a.accumulated_depreciation,
            "book_value": a.book_value if a.status != "DISPOSED" else ZERO,
        }
        for a in assets
    ]
    return ReportResult(
        "asset-register",
        "Fixed asset register",
        [
            Column("asset_tag", "Tag"),
            Column("name", "Asset"),
            Column("category", "Category"),
            Column("status", "Status"),
            Column("purchase_date", "Purchased", "date"),
            Column("cost", "Cost", "money"),
            Column("accumulated_depreciation", "Accumulated depreciation", "money"),
            Column("book_value", "Book value", "money"),
        ],
        rows,
        _sum(rows, "cost", "accumulated_depreciation", "book_value"),
        {"as_of": local_today()},
    )


REPORTS: dict[str, ReportDefinition] = {
    d.key: d
    for d in (
        ReportDefinition(
            "sales-summary",
            "Sales summary",
            "Issued invoices grouped by month, day or customer.",
            True,
            sales_summary,
            ("month", "day", "customer"),
        ),
        ReportDefinition(
            "sales-by-product",
            "Sales by product",
            "Quantity, revenue, cost of sales and margin per product.",
            True,
            sales_by_product,
        ),
        ReportDefinition(
            "profit-and-loss",
            "Profit and loss",
            "Revenue less cost of sales and approved expenses.",
            True,
            profit_and_loss,
        ),
        ReportDefinition(
            "vat-summary",
            "VAT summary",
            "Output VAT on sales against input VAT on bills and expenses.",
            True,
            vat_summary,
        ),
        ReportDefinition(
            "purchases-by-supplier",
            "Purchases by supplier",
            "Approved purchase orders: ordered and received value.",
            True,
            purchases_by_supplier,
        ),
        ReportDefinition(
            "expenses-by-category",
            "Expenses by category",
            "Approved expenses per category.",
            True,
            expenses_by_category,
        ),
        ReportDefinition(
            "inventory-valuation",
            "Inventory valuation",
            "Stock on hand at weighted average cost, as of today.",
            False,
            inventory_valuation,
        ),
        ReportDefinition(
            "stock-movements",
            "Stock movements",
            "Quantity in and out per product for the period.",
            True,
            stock_movement_summary,
        ),
        ReportDefinition(
            "receivables-aging",
            "Receivables aging",
            "Customer balances by days overdue, as of today.",
            False,
            receivables_aging,
        ),
        ReportDefinition(
            "payables-aging",
            "Payables aging",
            "Supplier balances by days overdue, as of today.",
            False,
            payables_aging,
        ),
        ReportDefinition(
            "headcount",
            "Headcount",
            "Employees by department or branch and status.",
            False,
            headcount,
            ("department", "branch"),
        ),
        ReportDefinition(
            "asset-register",
            "Asset register",
            "Fixed assets with cost, depreciation and book value.",
            False,
            asset_register,
        ),
    )
}


def run_report(db: Session, key: str, params: ReportParams) -> ReportResult:
    definition = REPORTS.get(key)
    if definition is None:
        raise NotFoundError(f"Unknown report '{key}'")
    if params.date_to < params.date_from:
        raise BadRequestError("date_to cannot be before date_from")
    if params.group_by and params.group_by not in definition.group_by_options:
        options = ", ".join(definition.group_by_options) or "none"
        raise BadRequestError(f"Invalid group_by for this report. Options: {options}")
    result = definition.builder(db, params)
    result.parameters = {
        **(
            {"date_from": params.date_from, "date_to": params.date_to}
            if definition.uses_dates
            else {}
        ),
        **({"group_by": params.group_by} if params.group_by else {}),
        **result.parameters,
    }
    return result


# ---------------------------------------------------------------------- export
def _plain(value: Any) -> Any:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, date):
        return value.isoformat()
    return value


def to_csv(result: ReportResult) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([c.label for c in result.columns])
    for row in result.rows:
        writer.writerow([_plain(row.get(c.key, "")) for c in result.columns])
    if result.totals:
        writer.writerow(
            [
                "Total" if i == 0 else _plain(result.totals.get(c.key, ""))
                for i, c in enumerate(result.columns)
            ]
        )
    # UTF-8 with BOM so Excel opens accented names correctly.
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(result: ReportResult) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = result.title[:31]
    ws.append([result.title])
    ws["A1"].font = Font(bold=True, size=14)
    for key, value in result.parameters.items():
        ws.append([key.replace("_", " ").capitalize(), str(_plain(value))])
    ws.append([])
    ws.append([c.label for c in result.columns])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    formats = {"money": "#,##0.00", "quantity": "#,##0.###", "percent": "0.00", "number": "0"}
    for row in result.rows:
        ws.append([_plain(row.get(c.key, "")) for c in result.columns])
        for cell, column in zip(ws[ws.max_row], result.columns, strict=False):
            if column.kind in formats:
                cell.number_format = formats[column.kind]
    if result.totals:
        ws.append(
            [
                "Total" if i == 0 else _plain(result.totals.get(c.key, ""))
                for i, c in enumerate(result.columns)
            ]
        )
        for cell, column in zip(ws[ws.max_row], result.columns, strict=False):
            cell.font = Font(bold=True)
            if column.kind in formats:
                cell.number_format = formats[column.kind]
    for i, column in enumerate(result.columns, start=1):
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = max(
            12, len(column.label) + 4
        )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def month_start(today: date | None = None) -> date:
    today = today or local_today()
    return today.replace(day=1)
